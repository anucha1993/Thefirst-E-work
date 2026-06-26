#!/usr/bin/env python3
"""
Probe: หลังจาก Step 2.1 สำเร็จ - ดู elements ที่อยู่บน page
ตามหา:
1. Edit button / "แก้ไข" link
2. Address fields
3. Permit number field/display
"""

from pathlib import Path
from playwright.sync_api import sync_playwright
import json
import time

LOGIN = "Vararat@mailforeign.com"
PASSWORD = "Q1w2e3r4@!"
EXCEL_FILE = Path("from_bt44.xlxs.xlsx")

def main():
    with sync_playwright() as pw:
        browser = pw.chromium.launch(args=["--ignore-certificate-errors", "--start-maximized"])
        ctx = browser.new_context(
            locale="th-TH",
            ignore_https_errors=True,
            viewport={"width": 1920, "height": 1080}
        )
        page = ctx.new_page()
        
        try:
            # === LOGIN ===
            page.goto("https://eworkpermit.doe.go.th/", wait_until="load")
            page.wait_for_load_state("networkidle", timeout=15000)
            time.sleep(2)
            
            page.wait_for_selector("input[name='username']", timeout=10000)
            page.fill("input[name='username']", LOGIN)
            page.fill("input[name='password']", PASSWORD)
            page.click("button:has-text('เข้าสู่ระบบ')")
            page.wait_for_load_state("networkidle", timeout=15000)
            time.sleep(2)
            
            # === OPEN BT44 ===
            page.evaluate("openCity('tab_CHANGE_REQ')")
            time.sleep(1)
            page.wait_for_selector("#CHANGE_EMPLOYER", timeout=5000)
            page.click("#CHANGE_EMPLOYER")
            page.wait_for_load_state("networkidle")
            time.sleep(1)
            
            # === FILL STEP 1 ===
            # Find search inputs
            inputs = page.query_selector_all("input[type='text']")
            print(f"Found {len(inputs)} text inputs on page")
            
            # Get search value from Excel
            from openpyxl import load_workbook
            wb = load_workbook(str(EXCEL_FILE), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                print("No data in Excel")
                return
            
            # Get first data row
            headers = rows[0]
            data_row = rows[1]
            data = dict(zip(headers, data_row))
            
            print(f"\nData: {data.get('ชื่อผู้กระทำการแทน')} | {data.get('หมายเลขอ้างอิง')}")
            
            # Search for alien
            search_fields = page.query_selector_all("input[type='text']")
            if search_fields:
                search_fields[0].fill(str(data.get("หมายเลขอ้างอิง", "")).strip())
                page.wait_for_timeout(500)
            
            page.click("button:has-text('ค้นหา')")
            page.wait_for_load_state("networkidle")
            time.sleep(2)
            
            # Check if results appeared
            results_table = page.query_selector("table")
            if results_table:
                rows = page.query_selector_all("tbody tr")
                print(f"Found {len(rows)} results")
                
                # Click first result
                if rows:
                    rows[0].click()
                    page.wait_for_load_state("networkidle")
                    time.sleep(1)
            
            # === STEP 2.1: Check consent ===
            print("\n=== STEP 2.1: Checking consent ===")
            checkbox = page.query_selector("#check_truth")
            if checkbox:
                checkbox.click()
                page.wait_for_timeout(300)
            
            # Click ถัดไป
            next_buttons = page.query_selector_all("button, a")
            for btn in next_buttons:
                txt = page.evaluate("el => el.textContent", btn)
                if "ถัดไป" in txt:
                    btn.click()
                    break
            
            page.wait_for_timeout(2000)
            
            # === NOW PROBE ELEMENTS ===
            print("\n=== AFTER STEP 2.1 - PROBING ELEMENTS ===")
            
            # 1. All buttons and links with text containing keywords
            page.screenshot(path="step2_after_consent.png", full_page=True)
            print("Screenshot saved: step2_after_consent.png\n")
            
            all_buttons = page.evaluate(r"""() => {
                const buttons = document.querySelectorAll('button, a, [role="button"]');
                const result = [];
                for (const btn of buttons) {
                    const txt = (btn.textContent || '').trim().substring(0, 50);
                    const visible = btn.offsetParent !== null;
                    if (txt && visible) {
                        result.push({
                            tag: btn.tagName,
                            text: txt,
                            id: btn.id,
                            class: btn.className.substring(0, 100),
                            visible: visible
                        });
                    }
                }
                return result;
            }""")
            
            print("All visible buttons/links:")
            for btn in all_buttons:
                if len(btn["text"]) > 2:
                    print(f"  - {btn['tag']}: '{btn['text']}' [id={btn['id']}, class={btn['class']}]")
            
            # 2. Look for edit buttons specifically
            print("\n\nSearching for 'แก้ไข' buttons:")
            edit_buttons = page.evaluate(r"""() => {
                const buttons = document.querySelectorAll('button, a, [role="button"], span, div');
                const result = [];
                for (const btn of buttons) {
                    const txt = (btn.textContent || '').trim();
                    if (/แก้ไข|EDIT|Edit|modify/i.test(txt)) {
                        const visible = btn.offsetParent !== null;
                        result.push({
                            tag: btn.tagName,
                            text: txt.substring(0, 100),
                            visible: visible,
                            onclick: btn.onclick ? "yes" : "no",
                            classes: btn.className
                        });
                    }
                }
                return result;
            }""")
            
            if edit_buttons:
                for btn in edit_buttons:
                    print(f"  Found: {btn}")
            else:
                print("  No 'แก้ไข' buttons found!")
            
            # 3. Look for address section
            print("\n\nSearching for address section:")
            addr_info = page.evaluate(r"""() => {
                // Find sections with address text
                const sections = document.querySelectorAll('div, section, fieldset');
                const result = [];
                for (const sec of sections) {
                    const txt = (sec.textContent || '').substring(0, 200);
                    if (/ที่อยู่|address|เขต|จังหวัด/i.test(txt) && txt.length > 20) {
                        result.push(txt);
                    }
                }
                return result;
            }""")
            
            if addr_info:
                print(f"Found {len(addr_info)} address-related sections")
                for i, txt in enumerate(addr_info[:3]):
                    print(f"  [{i}] {txt[:100]}...")
            
            # 4. Look for permit number
            print("\n\nSearching for permit number:")
            permit_info = page.evaluate(r"""() => {
                const text = document.body.innerText;
                const match = text.match(/\d{12}/);
                if (match) {
                    return match[0];
                }
                
                // Try specific field
                const fields = document.querySelectorAll('input, span, p, div');
                for (const field of fields) {
                    const txt = (field.textContent || '').trim();
                    if (/\d{12}/.test(txt) && txt.length < 50) {
                        return txt;
                    }
                }
                return null;
            }""")
            
            if permit_info:
                print(f"Found permit-like number: {permit_info}")
            else:
                print("No permit number found!")
            
            print("\n✓ Probe complete")
            
        except Exception as e:
            print(f"Error: {e}")
        finally:
            page.screenshot(path="step2_final_state.png", full_page=True)
            print("Final state screenshot: step2_final_state.png")
            ctx.close()
            browser.close()

if __name__ == "__main__":
    main()
