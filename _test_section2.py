"""ทดสอบ section 2 (capture permit info) standalone — ใช้บัญชีที่ register สำเร็จแล้ว"""
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
import sys; sys.path.insert(0, ".")
from scrape_wa import _capture_alien_permit_info, _save_register_report, _make_basename

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[3])}
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")
BASE = _make_basename(r)
print("Basename:", BASE)

screenshot_dir = Path("screenshots/register"); screenshot_dir.mkdir(parents=True, exist_ok=True)

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors", "--start-maximized"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True, viewport={"width":1920,"height":1080})
    page = ctx.new_page()
    info = _capture_alien_permit_info(page, EMAIL, PWD, basename=BASE, screenshot_dir=screenshot_dir, log=print)
    ctx.close(); b.close()

print("\n=== PermitFields ===")
for k, v in (info.get("PermitFields") or {}).items():
    print(f"  {k}: {v}")
print("Screenshot:", info.get("PermitScreenshot"))
print("ProfileScreenshot:", info.get("ProfileScreenshot"))
print("Error:", info.get("PermitError"))

# save report
res = {**r, "Status": "SUCCESS", "AttemptedWith": "TaxID", **info,
       "ScreenshotSuccess": f"{BASE}_success.png", "AlertText": "", "Error": "",
       "ScreenshotAlert": "", "RowIndex": 2}
_save_register_report([res], Path("WA_register_report_test4.xlsx"), log=print)
