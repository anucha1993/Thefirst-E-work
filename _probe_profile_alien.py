"""probe https://eworkpermit.doe.go.th/profile_alien — ดู DOM ของ รหัสอ้างอิง / รหัสคนต่างด้าว / สถานะคำขอ"""
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
import sys; sys.path.insert(0, ".")
from scrape_wa import _switch_lang_th

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]
r = {hdr[i]: c.value for i, c in enumerate(ws[3])}  # row 2 (index 3 = SU LATT HTWE)
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")
print("Login as:", EMAIL)

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login/", wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(1500)
    page.locator("#User_typeLogin").select_option(value="1")
    page.wait_for_timeout(500)
    page.locator("#username_login").fill(EMAIL)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    page.wait_for_timeout(4000)
    print("after login url:", page.url)
    _switch_lang_th(page)
    page.goto("https://eworkpermit.doe.go.th/profile_alien", wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(4000)
    print("profile url:", page.url)
    _switch_lang_th(page)
    page.wait_for_timeout(2000)
    page.screenshot(path="screenshots/register/_probe_profile_alien.png", full_page=True)
    # save html
    html = page.content()
    open("_probe_profile_alien.html", "w", encoding="utf-8").write(html)
    body_text = page.evaluate("() => document.body.innerText")
    open("_probe_profile_alien.txt", "w", encoding="utf-8").write(body_text)
    print("body text head:")
    print(body_text[:2000])
    print("\n--- search keywords ---")
    for kw in ["รหัสอ้างอิง", "รหัสคนต่างด้าว", "อยู่ระหว่าง", "ใบอนุญาต"]:
        for line in body_text.splitlines():
            if kw in line:
                print(f"  [{kw}] {line!r}")
    page.wait_for_timeout(3000)
    ctx.close(); b.close()
