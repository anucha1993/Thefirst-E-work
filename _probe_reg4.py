import sys, json, re
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; row1 = {hdr[i]: c.value for i, c in enumerate(ws[2])}
TAX = str(row1.get("TaxID") or ""); EMAIL=str(row1.get("Email") or ""); PWD=str(row1.get("Password") or "")

def dump_visible(page, max_btns=40):
    return page.evaluate(r"""(maxBtns) => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const radios = Array.from(document.querySelectorAll('input[type=radio]')).filter(visible).map(e => ({id:e.id, name:e.name, value:e.value, checked:e.checked, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,80) : ''}));
      const checkboxes = Array.from(document.querySelectorAll('input[type=checkbox]')).filter(visible).map(e => ({id:e.id, name:e.name, checked:e.checked, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,80) : (e.parentElement||{}).textContent?.trim().slice(0,80)}));
      const inputs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=radio]):not([type=checkbox]):not([type=image])')).filter(visible).map(e => ({id:e.id, name:e.name, type:e.type, placeholder:e.placeholder, value:e.value, readonly: e.readOnly, disabled: e.disabled}));
      const textareas = Array.from(document.querySelectorAll('textarea')).filter(visible).map(e => ({id:e.id, name:e.name, placeholder:e.placeholder}));
      const selects = Array.from(document.querySelectorAll('select')).filter(visible).map(e => ({id:e.id, name:e.name, opts: Array.from(e.options).map(o => `${o.value}::${o.text.trim()}`).slice(0,15)}));
      const buttons = Array.from(document.querySelectorAll('button, a.btn')).filter(visible).map(e => ({tag:e.tagName, id:e.id, txt:(e.textContent||'').trim().slice(0,50), onclick:(e.getAttribute('onclick')||'').slice(0,80)})).slice(0, maxBtns);
      // Also peek any prominent headings
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,legend')).filter(visible).map(e => e.textContent.trim().slice(0,100)).slice(0,10);
      return {url: location.href, heads, radios, checkboxes, inputs, textareas, selects, buttons};
    }""", max_btns)

def collect_alerts(page):
    return page.evaluate(r"""() => {
      const out = [];
      document.querySelectorAll('.swal2-popup, .swal2-container').forEach(el => { if (el.offsetParent === null) return; const t=el.querySelector('.swal2-title'); const c=el.querySelector('.swal2-html-container, .swal2-content'); out.push({kind:'swal', title:t?t.textContent.trim():'', content:c?c.textContent.trim():el.textContent.trim().slice(0,300)}); });
      document.querySelectorAll('.modal.show, .modal[style*="display: block"]').forEach(el => { const t=el.querySelector('.modal-title'); const b=el.querySelector('.modal-body'); out.push({kind:'modal', id:el.id, title:t?t.textContent.trim():'', body:b?b.textContent.trim().slice(0,400):''}); });
      return out;
    }""")

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login/Register", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeRegister").select_option(value="1"); page.wait_for_timeout(700)
    page.locator("input#alienOption1").check(force=True); page.wait_for_timeout(400)
    page.locator("#typeAlienUser").select_option(value="1"); page.wait_for_timeout(400)
    page.locator("#inputAlien1").fill(TAX); page.wait_for_timeout(300)
    page.locator("#validate_register").click(); page.wait_for_timeout(3500)
    page.screenshot(path="screenshots/reg_step_a_terms.png", full_page=True)
    print("=== AFTER submit (Register/Alien) ==="); print(json.dumps(dump_visible(page), ensure_ascii=False, indent=2))

    # คลิก ถัดไป (terms)
    try:
        page.locator("#gonext_0").click(); page.wait_for_timeout(2500)
    except Exception as e:
        print("gonext_0 click err:", e)
    page.screenshot(path="screenshots/reg_step_b_after_next.png", full_page=True)
    print("\n=== AFTER ถัดไป #gonext_0 ==="); print(json.dumps(dump_visible(page), ensure_ascii=False, indent=2))
    print("ALERTS:", json.dumps(collect_alerts(page), ensure_ascii=False))

    ctx.close(); b.close()
