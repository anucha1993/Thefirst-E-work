import json, re
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[3])}
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeLogin").select_option(value="1"); page.wait_for_timeout(1500)
    page.locator("#username_login").fill(EMAIL)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    page.wait_for_timeout(6000)
    try:
        page.locator("#sltLang").first.select_option(value="th"); page.wait_for_timeout(2500)
    except Exception as e:
        print("lang err:", e)

    page.goto("https://eworkpermit.doe.go.th/Profile/ProfileDigitalWorkPermit", wait_until="domcontentloaded")
    page.wait_for_timeout(4000)
    try:
        page.locator("#sltLang").first.select_option(value="th"); page.wait_for_timeout(2000)
    except Exception:
        pass
    page.screenshot(path="screenshots/digital_work_permit.png", full_page=True)

    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      // ดึง label-value pairs
      const rows = [];
      document.querySelectorAll('label, .form-label, .col-form-label, dt, th, .text-label').forEach(l => {
        if (!visible(l)) return;
        const txt = l.textContent.trim().slice(0,80);
        if (!txt) return;
        // หา value ถัดไป
        let v = '';
        const sib = l.nextElementSibling;
        if (sib) v = sib.textContent.trim().slice(0,200);
        const par = l.parentElement;
        if (!v && par) {
          const cands = par.querySelectorAll('span, div, dd, td, input');
          if (cands.length) {
            for (const c of cands) {
              if (c === l) continue;
              const t = (c.value || c.textContent || '').trim();
              if (t && t !== txt) { v = t.slice(0,200); break; }
            }
          }
        }
        rows.push({label: txt, value: v});
      });
      // Tabs/headings/section
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,h5,h6,legend,.card-title')).filter(visible).map(e => e.textContent.trim().slice(0,150)).slice(0,30);
      // Tables
      const tables = [];
      document.querySelectorAll('table').forEach(t => {
        if (!visible(t)) return;
        const headers = Array.from(t.querySelectorAll('thead th, thead td')).map(e => e.textContent.trim().slice(0,60));
        const body_rows = Array.from(t.querySelectorAll('tbody tr')).slice(0,30).map(tr => Array.from(tr.querySelectorAll('th,td')).map(c => c.textContent.trim().slice(0,150)));
        tables.push({headers, body_rows});
      });
      return {url: location.href, title: document.title, heads, rows: rows.slice(0,80), tables, body_text: document.body.innerText.slice(0, 3000)};
    }""")
    print(json.dumps(info, ensure_ascii=False, indent=2))

    ctx.close(); b.close()
