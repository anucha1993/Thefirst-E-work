"""ทดสอบ run_namelist_alien: ดึง 100 rows แรกจาก NameListAlien"""
import os, sys
from pathlib import Path
from dotenv import load_dotenv

sys.path.insert(0, str(Path(__file__).parent))
load_dotenv()

from scrape_wa import run_namelist_alien

cfg = {
    "username": os.getenv("EWP_USERNAME", ""),
    "password": os.getenv("EWP_PASSWORD", ""),
    "user_type": os.getenv("EWP_USER_TYPE", "บริษัทนำเข้า (บนจ.)"),
    "method": os.getenv("EWP_LOGIN_METHOD", "E-Workpermit"),
    "headless": False,
    "hide_window": False,
}
print(f"[i] ใช้บัญชี: {cfg['username']} ({cfg['user_type']})")

out_path = Path("reports/_test_namelist_alien.xlsx")
out_path.parent.mkdir(parents=True, exist_ok=True)

try:
    count, path = run_namelist_alien(
        cfg, out_path,
        form_type="MT_63_2_3103_RENEWAL",
        limit=100,  # ทดสอบ 100 rows แรก
        log=print,
    )
except Exception as e:
    import traceback; traceback.print_exc()
    sys.exit(2)

print(f"\n[✓] เสร็จ — count = {count}, report = {path}")

# แสดง 3 rows แรก
from openpyxl import load_workbook
wb = load_workbook(path, data_only=True)
ws = wb.active
print(f"\n[Report] {ws.max_row - 1} rows:")
headers = [c.value for c in ws[1]]
print("  headers:", headers)
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True), 1):
    print(f"  row {i}:")
    for h, v in zip(headers, row):
        val = str(v)[:80] if v is not None else "-"
        print(f"    {h}: {val}")
