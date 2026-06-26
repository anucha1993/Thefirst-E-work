#!/usr/bin/env python3
"""Debug probe: Check page state after Step 1 (search modal closes)"""
from playwright.sync_api import sync_playwright
import dotenv
import json

cfg = dotenv.dotenv_values(".env")

with sync_playwright() as p:
    browser = p.chromium.launch(
        headless=False,
        args=["--ignore-certificate-errors", "--start-maximized"]
    )
    ctx = browser.new_context(locale="th-TH", viewport={"width": 1920, "height": 1080}, ignore_https_errors=True)
    page = ctx.new_page()

    try:
        from scrape_wa import login, _open_bt44_form, _bt44_fill_search_one
        from pathlib import Path
        
        login(page, {
            "username": cfg.get("EWP_USERNAME") or cfg.get("username"),
            "password": cfg.get("EWP_PASSWORD") or cfg.get("password"),
            "user_type": cfg.get("EWP_USER_TYPE") or cfg.get("user_type"),
            "method": cfg.get("EWP_LOGIN_METHOD") or cfg.get("method", "E-Workpermit"),
        })
        page.wait_for_timeout(1500)
        
        _open_bt44_form(page)
        
        # Read test data
        from openpyxl import load_workbook
        wb = load_workbook("from_bt44.xlxs.xlsx", data_only=True)
        ws = wb.active
        row_data = list(ws.iter_rows(min_row=2, values_only=True))[0]
        rec = {
            "seq": "1",
            "prefix": row_data[1],
            "alien_ref": row_data[2],
            "name": row_data[3],
            "nationality": row_data[4],
            "sex": row_data[5],
            "birthdate": "25/11/2005",
            "username": "",
            "row_index": 1,
            "workpermit_no": str(row_data[7]),
        }
        
        ss_dir = Path("debug_screenshots")
        ss_dir.mkdir(exist_ok=True)
        
        # Run Step 1
        print("[Step 1] Running search...")
        res = _bt44_fill_search_one(page, rec, ss_dir)
        print(f"[Step 1] Result: {res.get('status')} - {res.get('note')}")
        
        page.wait_for_timeout(1500)
        page.screenshot(path=str(ss_dir / "after_step1.png"), full_page=True)
        
        # Check page state
        page_state = page.evaluate(r"""() => {
            const state = {};
            
            // Check for buttons containing "แก้ไข"
            state.edit_buttons = Array.from(document.querySelectorAll('button,a'))
                .filter(b => /แก้ไข|EDIT/i.test((b.textContent || '').trim()))
                .map(b => ({
                    text: b.textContent.trim().slice(0, 30),
                    visible: b.offsetParent !== null,
                    class: b.className
                }));
            
            // Check for buttons containing "ถัดไป"
            state.next_buttons = Array.from(document.querySelectorAll('button,a'))
                .filter(b => /ถัดไป|NEXT/i.test((b.textContent || '').trim()))
                .map(b => ({
                    text: b.textContent.trim().slice(0, 30),
                    visible: b.offsetParent !== null
                }));
            
            // Check for sections/forms
            state.sections = Array.from(document.querySelectorAll('section, fieldset, div[class*="section"], div[class*="form"]'))
                .slice(0, 10)
                .map(s => ({
                    tag: s.tagName.toLowerCase(),
                    class: s.className.slice(0, 50),
                    text: s.textContent.trim().slice(0, 100)
                }));
            
            // Check for work permit field
            state.permit_fields = Array.from(document.querySelectorAll('input, select, span[id*="permit"], input[id*="permit"]'))
                .map(f => ({
                    tag: f.tagName.toLowerCase(),
                    id: f.id || f.name || '',
                    value: f.value || f.textContent.slice(0, 20),
                    text: f.textContent ? f.textContent.trim().slice(0, 50) : ''
                }));
            
            return state;
        }""")
        
        print("\n=== PAGE STATE AFTER STEP 1 ===")
        print(json.dumps(page_state, indent=2, ensure_ascii=False))
        
        # Get page title/heading
        title = page.evaluate(r"""() => {
            return document.title;
        }""")
        print(f"\nPage title: {title}")
        
    finally:
        browser.close()
