import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[2])}
TAX = str(r.get("TaxID") or ""); EMAIL=str(r.get("Email") or ""); PWD=str(r.get("Password") or "")

def dump(page, label, path):
    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const inputs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=image])')).filter(visible).map(e => ({id:e.id, name:e.name, type:e.type, value:e.value, checked:e.checked}));
      const buttons = Array.from(document.querySelectorAll('button, a.btn, a')).filter(visible).map(e => ({tag:e.tagName, id:e.id, href:e.getAttribute('href')||'', txt:(e.textContent||'').trim().slice(0,60), onclick:(e.getAttribute('onclick')||'').slice(0,80), disabled:e.disabled})).slice(0,40);
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,h5,legend,.h2,.h3')).filter(visible).map(e => e.textContent.trim().slice(0,150)).slice(0,15);
      const bodyText = document.body.innerText.slice(0, 1500);
      return {url:location.href, heads, inputs, buttons, bodyText};
    }""")
    print(f"\n=== {label} ===\n" + json.dumps(info, ensure_ascii=False, indent=2))
    page.screenshot(path=path, full_page=True)

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login/Register", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeRegister").select_option(value="1"); page.wait_for_timeout(400)
    page.locator("input#alienOption1").check(force=True); page.wait_for_timeout(300)
    page.locator("#typeAlienUser").select_option(value="1"); page.wait_for_timeout(300)
    page.locator("#inputAlien1").fill(TAX)
    page.locator("#validate_register").click(); page.wait_for_timeout(3500)
    page.locator("#checkbox_em").evaluate("e => e.click()"); page.wait_for_timeout(500)
    page.locator("#gonext_0").click(); page.wait_for_timeout(3500)
    page.locator("#educationLevel2").select_option(value="8"); page.wait_for_timeout(300)
    page.locator("#workExperien_Alien2").fill("2 ปี")
    page.locator("#validateFormAlien").click(); page.wait_for_timeout(3500)
    page.locator("#summaryInfoAlien").click(); page.wait_for_timeout(3500)
    page.locator("#setPassword").fill(PWD)
    page.locator("#confirmPassword").fill(PWD); page.wait_for_timeout(400)
    page.locator("#validateSetAlienPassword").click(); page.wait_for_timeout(3500)
    # หน้า OTP/Email - skip OTP + กรอก email
    page.locator("#otp_skip").evaluate("e => e.click()"); page.wait_for_timeout(300)
    page.locator("#user_email").fill(EMAIL); page.wait_for_timeout(300)
    page.locator("#validateEmailOTPAlien").click(); page.wait_for_timeout(5000)
    dump(page, "STEP F after email submit (final)", "screenshots/reg_step_final.png")
    ctx.close(); b.close()
