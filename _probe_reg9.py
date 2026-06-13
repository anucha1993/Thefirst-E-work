import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[2])}
TAX = str(r.get("TaxID") or ""); EMAIL=str(r.get("Email") or ""); PWD=str(r.get("Password") or "")

def dump(page, label, path):
    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const radios = Array.from(document.querySelectorAll('input[type=radio]')).filter(visible).map(e => ({id:e.id, name:e.name, value:e.value, checked:e.checked, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,90) : ''}));
      const checkboxes = Array.from(document.querySelectorAll('input[type=checkbox]')).filter(visible).map(e => ({id:e.id, name:e.name, checked:e.checked, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,90) : (e.parentElement||{}).textContent?.trim().slice(0,90)}));
      const inputs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=radio]):not([type=checkbox]):not([type=image])')).filter(visible).map(e => ({id:e.id, name:e.name, type:e.type, placeholder:e.placeholder, value:e.value}));
      const selects = Array.from(document.querySelectorAll('select')).filter(visible).map(e => ({id:e.id, name:e.name, opts: Array.from(e.options).map(o => `${o.value}::${o.text.trim()}`).slice(0,12)}));
      const buttons = Array.from(document.querySelectorAll('button, a.btn')).filter(visible).map(e => ({id:e.id, txt:(e.textContent||'').trim().slice(0,50), onclick:(e.getAttribute('onclick')||'').slice(0,80), disabled:e.disabled})).slice(0,30);
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,h5,legend,.h2,.h3')).filter(visible).map(e => e.textContent.trim().slice(0,120)).slice(0,15);
      return {url:location.href, heads, radios, checkboxes, inputs, selects, buttons};
    }""")
    print(f"\n=== {label} ===\n" + json.dumps(info, ensure_ascii=False, indent=2))
    page.screenshot(path=path, full_page=True)

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login/Register", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeRegister").select_option(value="1"); page.wait_for_timeout(400)
    page.locator("input#alienOption1").check(force=True); page.wait_for_timeout(300)
    page.locator("#typeAlienUser").select_option(value="1"); page.wait_for_timeout(300)
    page.locator("#inputAlien1").fill(TAX)
    page.locator("#validate_register").click(); page.wait_for_timeout(3500)
    page.locator("#checkbox_em").evaluate("e => e.click()"); page.wait_for_timeout(500)
    page.locator("#gonext_0").click(); page.wait_for_timeout(3500)
    page.locator("#educationLevel2").select_option(value="8"); page.wait_for_timeout(300)
    page.locator("#workExperien_Alien2").fill("2 ปี")
    page.locator("#validateFormAlien").click(); page.wait_for_timeout(3500)
    page.locator("#summaryInfoAlien").click(); page.wait_for_timeout(3500)
    page.locator("#setPassword").fill(PWD)
    page.locator("#confirmPassword").fill(PWD); page.wait_for_timeout(400)
    page.locator("#validateSetAlienPassword").click(); page.wait_for_timeout(3500)
    dump(page, "STEP E after password", "screenshots/reg_step_e.png")
    ctx.close(); b.close()
