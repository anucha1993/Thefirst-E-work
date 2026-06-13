import sys, json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True)
ws = wb.active
hdr = [c.value for c in ws[1]]
row1 = {hdr[i]: c.value for i, c in enumerate(ws[2])}
print("ROW1:", row1)
TAX = str(row1.get("TaxID") or "")
PP  = str(row1.get("PassportNo.") or "")

def collect_alerts(page):
    """หา SweetAlert/modal alert ที่ visible"""
    return page.evaluate(r"""() => {
      const out = [];
      // SweetAlert2
      document.querySelectorAll('.swal2-popup, .swal2-container').forEach(el => {
        if (el.offsetParent === null) return;
        const title = el.querySelector('.swal2-title'); const content = el.querySelector('.swal2-html-container, #swal2-content, .swal2-content');
        out.push({kind:'swal', title: title ? title.textContent.trim() : '', content: content ? content.textContent.trim() : el.textContent.trim().slice(0,300)});
      });
      // BS modal
      document.querySelectorAll('.modal.show, .modal[style*="display: block"]').forEach(el => {
        const title = el.querySelector('.modal-title'); const body = el.querySelector('.modal-body');
        out.push({kind:'modal', id: el.id, title: title ? title.textContent.trim() : '', body: body ? body.textContent.trim().slice(0,400) : ''});
      });
      return out;
    }""")

def dump_visible(page):
    return page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const radios = Array.from(document.querySelectorAll('input[type=radio]')).filter(visible).map(e => ({id:e.id, name:e.name, value:e.value, checked:e.checked, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,60) : ''}));
      const checkboxes = Array.from(document.querySelectorAll('input[type=checkbox]')).filter(visible).map(e => ({id:e.id, name:e.name, lbl: e.id ? (document.querySelector(`label[for="${e.id}"]`)||{}).textContent?.trim().slice(0,60) : ''}));
      const inputs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=radio]):not([type=checkbox]):not([type=image])')).filter(visible).map(e => ({id:e.id, name:e.name, type:e.type, placeholder:e.placeholder}));
      const textareas = Array.from(document.querySelectorAll('textarea')).filter(visible).map(e => ({id:e.id, name:e.name, placeholder:e.placeholder}));
      const selects = Array.from(document.querySelectorAll('select')).filter(visible).map(e => ({id:e.id, name:e.name, opts: Array.from(e.options).map(o => `${o.value}::${o.text.trim()}`).slice(0,12)}));
      const buttons = Array.from(document.querySelectorAll('button, a.btn')).filter(visible).map(e => ({tag:e.tagName, id:e.id, txt:(e.textContent||'').trim().slice(0,50)})).slice(0,30);
      return {url: location.href, radios, checkboxes, inputs, textareas, selects, buttons};
    }""")

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()

    page.goto("https://eworkpermit.doe.go.th/Login/Register", wait_until="domcontentloaded")
    page.wait_for_timeout(2500)
    page.locator("#User_typeRegister").select_option(value="1"); page.wait_for_timeout(800)
    page.locator("input#alienOption1").check(force=True); page.wait_for_timeout(500)
    page.locator("#typeAlienUser").select_option(value="1"); page.wait_for_timeout(500)
    page.locator("#inputAlien1").fill(TAX); page.wait_for_timeout(300)
    page.screenshot(path="screenshots/reg_before_submit.png", full_page=True)
    print(f"\n=== submit ด้วย TaxID={TAX} ===")
    page.locator("#validate_register").click()
    page.wait_for_timeout(3500)
    alerts = collect_alerts(page)
    print("ALERTS after TaxID submit:", json.dumps(alerts, ensure_ascii=False, indent=2))
    page.screenshot(path="screenshots/reg_after_taxid.png", full_page=True)
    print("\nVISIBLE STATE after TaxID:"); print(json.dumps(dump_visible(page), ensure_ascii=False, indent=2))
    ctx.close(); b.close()
