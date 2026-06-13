import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[3])}
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")

def dump_full(page):
    return page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const items = Array.from(document.querySelectorAll('button, a, [onclick]')).filter(visible).map(e => ({
        tag: e.tagName, id: e.id, href: e.getAttribute('href')||'',
        txt: (e.textContent||'').trim().slice(0,100),
        onclick: (e.getAttribute('onclick')||'').slice(0,150),
        cls: (e.className||'').toString().slice(0,80),
      })).filter(b => b.txt || b.onclick).slice(0, 200);
      return {url: location.href, items};
    }""")

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeLogin").select_option(value="1"); page.wait_for_timeout(1500)
    page.locator("#username_login").fill(EMAIL)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    page.wait_for_timeout(6000)

    # สลับภาษาเป็นไทยก่อน (มี 2 ตัว — desktop + sticky)
    try:
        page.locator("#sltLang").first.select_option(value="th"); page.wait_for_timeout(2500)
    except Exception as e:
        print("lang switch err:", e)
    page.screenshot(path="screenshots/post_login_th.png", full_page=True)

    info = dump_full(page)
    print("URL:", info["url"])
    print("\n=== HOME PAGE - filter by อนุญาต/profile/openpage ===")
    for it in info["items"]:
        text_match = any(k in it.get("txt","") for k in ["อนุญาต", "ขอใบ", "ข้อมูล", "ข้อมูลส่วน", "ออกจาก"])
        oc_match = any(k in it.get("onclick","") for k in ["openprofile", "openpage", "openmenu", "Permit", "logout"])
        if text_match or oc_match:
            print(json.dumps(it, ensure_ascii=False))

    print("\n=== คลิก openprofile() ===")
    try:
        page.evaluate("openprofile()")
        page.wait_for_timeout(3500)
        print("URL after openprofile:", page.url)
        page.screenshot(path="screenshots/openprofile.png", full_page=True)
    except Exception as e:
        print("openprofile err:", e)

    info2 = dump_full(page)
    print("\n=== profile_alien items (ALL) ===")
    for it in info2["items"][:120]:
        print(json.dumps(it, ensure_ascii=False))
    page.screenshot(path="screenshots/profile_alien.png", full_page=True)

    ctx.close(); b.close()
