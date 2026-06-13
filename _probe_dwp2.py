import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[3])}
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeLogin").select_option(value="1"); page.wait_for_timeout(1500)
    page.locator("#username_login").fill(EMAIL)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    page.wait_for_timeout(6000)
    try:
        page.locator("#sltLang").first.select_option(value="th"); page.wait_for_timeout(2500)
    except Exception:
        pass

    page.goto("https://eworkpermit.doe.go.th/Profile/ProfileDigitalWorkPermit", wait_until="domcontentloaded")
    page.wait_for_timeout(4000)
    try:
        page.locator("#sltLang").first.select_option(value="th"); page.wait_for_timeout(2500)
    except Exception:
        pass

    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const sections = [];
      // หา block ภายใต้ section title
      document.querySelectorAll('.card, section, .panel, fieldset, .row').forEach(sec => {
        if (!visible(sec)) return;
        // หา label/value pair: pattern คือ <small>/.label-* คู่กับ <span>/<p>/<div>
        const pairs = [];
        sec.querySelectorAll('small, .label, .text-muted, .col-form-label, label').forEach(l => {
          if (!visible(l)) return;
          const txt = l.textContent.trim();
          if (!txt || txt.length > 80) return;
          // หาค่าใน sibling หรือ parent
          let v = '';
          let nxt = l.nextElementSibling;
          while (nxt && !v) {
            const t = nxt.textContent.trim();
            if (t && t !== txt && t.length < 400) { v = t; break; }
            nxt = nxt.nextElementSibling;
          }
          if (!v && l.parentElement) {
            const cands = l.parentElement.querySelectorAll('span, b, strong, p, div.value, h5, h6');
            for (const c of cands) {
              if (!visible(c) || c === l || c.contains(l) || l.contains(c)) continue;
              const t = c.textContent.trim();
              if (t && t !== txt && t.length < 400) { v = t; break; }
            }
          }
          if (v) pairs.push({label: txt, value: v});
        });
        if (pairs.length) sections.push({pairs: pairs.slice(0, 30)});
      });
      return {url: location.href, sections: sections.slice(0, 15)};
    }""")
    print(json.dumps(info, ensure_ascii=False, indent=2))
    ctx.close(); b.close()
