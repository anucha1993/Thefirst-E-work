import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r2 = {hdr[i]: c.value for i, c in enumerate(ws[3])}
print("Row 2:", r2)
TAX = str(r2.get("TaxID") or ""); PWD = str(r2.get("Password") or "")

def dump(page, label, path=None):
    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const inputs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=image])')).filter(visible).map(e => ({id:e.id, name:e.name, type:e.type, placeholder:e.placeholder, value:e.value}));
      const selects = Array.from(document.querySelectorAll('select')).filter(visible).map(e => ({id:e.id, name:e.name, opts: Array.from(e.options).map(o => `${o.value}::${o.text.trim()}`).slice(0,12)}));
      const buttons = Array.from(document.querySelectorAll('button, a')).filter(visible).map(e => ({tag:e.tagName, id:e.id, href:e.getAttribute('href')||'', txt:(e.textContent||'').trim().slice(0,80), onclick:(e.getAttribute('onclick')||'').slice(0,80)})).filter(b => b.txt).slice(0, 60);
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,h5,legend,.h2,.h3')).filter(visible).map(e => e.textContent.trim().slice(0,120)).slice(0,15);
      return {url: location.href, heads, inputs, selects, buttons};
    }""")
    print(f"\n=== {label} ===\n" + json.dumps(info, ensure_ascii=False, indent=2))
    if path: page.screenshot(path=path, full_page=True)

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    dump(page, "LOGIN page")

    # เลือกประเภทผู้ใช้งาน "คนต่างด้าว"
    page.locator("select").nth(1).select_option(label="คนต่างด้าว"); page.wait_for_timeout(1500)
    page.screenshot(path="screenshots/login_alien.png", full_page=True)
    dump(page, "after select คนต่างด้าว")

    # method ส่วนใหญ่ default = E-Workpermit. กรอก
    user_box = page.locator("#username_login:visible, #employer_login:visible, #agency_login:visible").first
    user_box.wait_for(state="visible", timeout=10000)
    user_box.fill(TAX)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    try:
        page.wait_for_url(lambda u: "/Login" not in u, timeout=30000)
    except Exception as e:
        print("URL wait err:", e)
    page.wait_for_timeout(3000)
    page.screenshot(path="screenshots/login_after.png", full_page=True)
    dump(page, "AFTER LOGIN", "screenshots/post_login.png")

    ctx.close(); b.close()
