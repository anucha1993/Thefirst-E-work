"""Probe: read Request_Receipt.xlsx + sample PAYMENT_RECEIPT PDFs to understand structure."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).parent
xlsx = ROOT / "Request_Receipt.xlsx"
print(f"=== {xlsx.name} ===")
wb = load_workbook(xlsx, read_only=True, data_only=True)
ws = wb.active
print(f"Sheet: {ws.title}, max_row={ws.max_row}, max_col={ws.max_column}")
for row in ws.iter_rows(min_row=1, max_row=min(6, ws.max_row), values_only=True):
    print(row)

# PDF samples
from pypdf import PdfReader
for f in ["Files/69125200693421_PAYMENT_RECEIPT_20260408183225.pdf",
          "Files/69125200694755_PAYMENT_RECEIPT_20260421205955.pdf"]:
    p = ROOT / f
    print(f"\n=== {p.name} ===")
    r = PdfReader(str(p))
    print(f"Pages: {len(r.pages)}")
    for i, page in enumerate(r.pages[:3], start=1):
        txt = page.extract_text() or ""
        print(f"\n--- Page {i} (first 1500 chars) ---")
        print(txt[:1500])
    if len(r.pages) > 3:
        print(f"\n--- last page (page {len(r.pages)}) ---")
        print((r.pages[-1].extract_text() or "")[:1500])
