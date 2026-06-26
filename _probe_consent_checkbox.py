#!/usr/bin/env python3
"""Debug: Find consent checkbox"""
from playwright.sync_api import sync_playwright
import dotenv
import json
from pathlib import Path

cfg = dotenv.dotenv_values(".env")

with sync_playwright() as p:
    browser = p.chromium.launch(headless=False, args=["--ignore-certificate-errors", "--start-maximized"])
    ctx = browser.new_context(locale="th-TH", viewport={"width": 1920, "height": 1080}, ignore_https_errors=True)
    page = ctx.new_page()

    try:
        from scrape_wa import login, _open_bt44_form, _bt44_fill_search_one
        from openpyxl import load_workbook
        
        login(page, {
            "username": cfg.get("EWP_USERNAME", ""),
            "password": cfg.get("EWP_PASSWORD", ""),
            "user_type": cfg.get("EWP_USER_TYPE", ""),
            "method": cfg.get("EWP_LOGIN_METHOD", "E-Workpermit"),
        })
        page.wait_for_timeout(1500)
        
        _open_bt44_form(page)
        
        # Read test data
        wb = load_workbook("from_bt44.xlxs.xlsx", data_only=True)
        ws = wb.active
        row_data = list(ws.iter_rows(min_row=2, values_only=True))[0]
        rec = {
            "seq": "1", "prefix": row_data[1], "alien_ref": row_data[2], "name": row_data[3],
            "nationality": row_data[4], "sex": row_data[5], "birthdate": "25/11/2005",
            "username": "", "row_index": 1, "workpermit_no": str(row_data[7]),
        }
        
        ss_dir = Path("debug_screenshots")
        ss_dir.mkdir(exist_ok=True)
        
        # Run Step 1
        print("[Step 1] Running search...")
        res = _bt44_fill_search_one(page, rec, ss_dir)
        print(f"[Step 1] Result: {res.get('status')}")
        page.wait_for_timeout(1500)
        
        # Close any modals
        page.evaluate(r"""() => {
            const modals = document.querySelectorAll('.modal, .swal2-popup, [role="dialog"]');
            for (const m of modals) {
                if (m.offsetParent !== null) {
                    const closeBtn = m.querySelector('.close, .btn-close, .swal2-close, .btn[onclick*="close"]');
                    if (closeBtn) closeBtn.click();
                }
            }
        }""")
        page.wait_for_timeout(800)
        
        # Now look for checkboxes
        checkboxes_info = page.evaluate(r"""() => {
            const checkboxes = [];
            document.querySelectorAll('input[type="checkbox"]').forEach((chk, idx) => {
                const label = chk.closest('label') || document.querySelector(`label[for="${chk.id}"]`);
                const parent = chk.closest('div[class*="form"], div[class*="group"]') || chk.parentElement;
                checkboxes.push({
                    idx: idx,
                    id: chk.id,
                    name: chk.name,
                    checked: chk.checked,
                    visible: chk.offsetParent !== null,
                    label_text: (label?.textContent || parent?.textContent || '').trim().slice(0, 200),
                    parent_class: parent?.className || '',
                });
            });
            return checkboxes;
        }""")
        
        print("\n=== CHECKBOXES FOUND ===")
        print(json.dumps(checkboxes_info, indent=2, ensure_ascii=False))
        
        # Try to find and check the consent checkbox
        print("\n[Attempting to check consent checkbox...]")
        checked = page.evaluate(r"""() => {
            const chk = Array.from(document.querySelectorAll('input[type="checkbox"]'))
              .find(c => {
                const txt = (c.closest('label')?.textContent || c.parentElement?.textContent || '').trim();
                console.log('Checkbox text:', txt.slice(0, 100));
                return /ข้าพเจ้าขอรับรอง/.test(txt);
              });
            if (chk) {
              console.log('Found consent checkbox');
              chk.checked = true;
              chk.dispatchEvent(new Event('change', { bubbles: true }));
              return true;
            }
            console.log('Consent checkbox not found');
            return false;
        }""")
        print(f"✓ Consent checkbox found and checked: {checked}")
        
        page.screenshot(path=str(ss_dir / "after_step2_1.png"), full_page=True)
        
    finally:
        browser.close()
