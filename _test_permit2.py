"""ทดสอบดึง permit_report แค่ 2 record (live) เพื่อยืนยันคอลัมน์ 'สถานะคำขอ' แบบทุกสถานะ
รัน:  python _test_permit2.py
- เบราว์เซอร์จะเปิดขึ้น → แก้ reCAPTCHA ให้ทันภายใน ~30 วิ แล้วปล่อยให้ทำงานต่อ
"""
from pathlib import Path

from openpyxl import load_workbook

from scrape_wa import load_config, run_permit_report

cfg = load_config()  # อ่าน username/password/user_type/method/headless จาก .env
# ใช้รายการคำขอ CHANGE_44_22 (รู้ว่ามีข้อมูล + ไทม์ไลน์สถานะครบ)
cfg["request_types"] = ["CHANGE_44_22"]
cfg["request_type"] = "CHANGE_44_22"
cfg["filter_status_ids"] = ["WP", "WCOSNA", "WA", "AP", "SS"]  # ทุกสถานะ กันตกหล่น
cfg["date_from"] = ""
cfg["date_to"] = ""
cfg["headless"] = False  # ต้องเห็นหน้าต่างเพื่อแก้ reCAPTCHA
cfg["save_every"] = 999

out = Path("reports/_test_permit2.xlsx")
count, path = run_permit_report(cfg, out, limit=2)

print(f"\n================ DONE: {count} rows -> {path} ================")
wb = load_workbook(path)
ws = wb.active
print("header:", [c.value for c in ws[1]])
for i in range(2, ws.max_row + 1):
    print(f"\n--- row {i - 1}: {ws.cell(i, 3).value} | {ws.cell(i, 1).value} | {ws.cell(i, 2).value} ---")
    print(ws.cell(i, 4).value)
