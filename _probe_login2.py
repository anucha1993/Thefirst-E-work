import json
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

wb = load_workbook("Exm.xlsx", data_only=True); ws = wb.active
hdr = [c.value for c in ws[1]]; r = {hdr[i]: c.value for i, c in enumerate(ws[3])}
EMAIL = str(r.get("Email") or ""); PWD = str(r.get("Password") or "")

def alerts(page):
    return page.evaluate(r"""() => {
      const out=[]; document.querySelectorAll('.swal2-popup, .swal2-container').forEach(el=>{ if(el.offsetParent===null) return; const t=el.querySelector('.swal2-title'),c=el.querySelector('.swal2-html-container, .swal2-content'); out.push({k:'swal',t:t?t.textContent.trim():'',c:c?c.textContent.trim():el.textContent.trim().slice(0,300)}); }); return out;
    }""")

def dump(page, label):
    info = page.evaluate(r"""() => {
      const visible = e => e.offsetParent !== null && getComputedStyle(e).visibility !== 'hidden';
      const buttons = Array.from(document.querySelectorAll('button, a')).filter(visible).map(e => ({tag:e.tagName, id:e.id, href:e.getAttribute('href')||'', txt:(e.textContent||'').trim().slice(0,80), onclick:(e.getAttribute('onclick')||'').slice(0,80), cls:e.className.slice(0,80)})).filter(b => b.txt).slice(0, 80);
      const heads = Array.from(document.querySelectorAll('h1,h2,h3,h4,h5')).filter(visible).map(e => e.textContent.trim().slice(0,120)).slice(0,15);
      const navs = Array.from(document.querySelectorAll('nav, .navbar, .dropdown-menu, [class*="user-name"], [class*="profile"]')).filter(visible).slice(0,12).map(e => ({tag:e.tagName, id:e.id, cls:e.className.slice(0,80), txt:e.textContent.trim().slice(0,250)}));
      return {url: location.href, heads, buttons, navs};
    }""")
    print(f"\n=== {label} ===\n" + json.dumps(info, ensure_ascii=False, indent=2))

with sync_playwright() as pw:
    b = pw.chromium.launch(headless=False, args=["--ignore-certificate-errors"])
    ctx = b.new_context(locale="th-TH", ignore_https_errors=True)
    page = ctx.new_page()
    page.goto("https://eworkpermit.doe.go.th/Login", wait_until="domcontentloaded"); page.wait_for_timeout(2500)
    page.locator("#User_typeLogin").select_option(value="1"); page.wait_for_timeout(1500)
    page.locator("#username_login").fill(EMAIL)
    page.locator("#password_login").fill(PWD)
    page.locator("#validate_login").click()
    page.wait_for_timeout(6000)
    print("ALERTS:", json.dumps(alerts(page), ensure_ascii=False))
    print("URL:", page.url)
    page.screenshot(path="screenshots/login2_after.png", full_page=True)
    dump(page, "POST LOGIN")
    ctx.close(); b.close()
